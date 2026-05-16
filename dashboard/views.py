import csv
from datetime import date, timedelta
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse
from django.utils import timezone
import zipfile
import io
from django.views.decorators.http import require_POST
from django.views.decorators.csrf import csrf_exempt
from .models import Room, EmployeeSession, PostureLog, WorkStatusLog, Warning, DailyStats, Notification, PersonalModel, FaceEncoding


@login_required
def index(request):
    """Ana sayfa - role gore yonlendir."""
    if request.user.is_yonetici() or request.user.is_takim_lideri():
        return redirect('dashboard:admin_dashboard')
    return redirect('dashboard:employee_panel')


@login_required
def admin_dashboard(request):
    """Yonetici canli izleme sayfasi."""
    if not request.user.is_yonetici() and not request.user.is_takim_lideri():
        return redirect('dashboard:employee_panel')

    rooms = request.user.get_accessible_rooms().filter(is_active=True)
    selected_room = None
    online_sessions = []

    room_code = request.GET.get('room')
    if room_code:
        selected_room = rooms.filter(code=room_code).first()
        if selected_room:
            online_sessions = EmployeeSession.objects.filter(
                room=selected_room, is_online=True
            ).select_related('user')

    return render(request, 'dashboard/admin_dashboard.html', {
        'rooms': rooms,
        'selected_room': selected_room,
        'online_sessions': online_sessions,
    })


@login_required
def room_manage(request):
    """Oda olusturma ve listeleme."""
    if not request.user.is_yonetici() and not request.user.is_takim_lideri():
        return redirect('dashboard:employee_panel')

    if request.method == 'POST' and request.user.can_create_room():
        name = request.POST.get('name', '').strip()
        if name:
            break_enabled = request.POST.get('break_reminder_enabled') == 'on'
            break_work = request.POST.get('break_work_minutes', '25')
            break_rest = request.POST.get('break_rest_minutes', '5')

            try:
                break_work = max(1, int(break_work))
                break_rest = max(1, int(break_rest))
            except (ValueError, TypeError):
                break_work = 25
                break_rest = 5

            Room.objects.create(
                name=name,
                created_by=request.user,
                break_reminder_enabled=break_enabled,
                break_work_minutes=break_work,
                break_rest_minutes=break_rest,
            )
            return redirect('dashboard:room_manage')

    rooms = request.user.get_accessible_rooms()
    return render(request, 'dashboard/room_manage.html', {
        'rooms': rooms,
        'can_create': request.user.can_create_room(),
        'can_delete': request.user.can_delete_room(),
    })


@login_required
def room_toggle(request, room_id):
    """Oda aktif/pasif degistir."""
    if not request.user.can_create_room():
        return redirect('dashboard:employee_panel')

    room = get_object_or_404(Room, id=room_id)
    if not request.user.is_super_admin() and room.created_by != request.user:
        return redirect('dashboard:room_manage')
    room.is_active = not room.is_active
    room.save()
    return redirect('dashboard:room_manage')


@login_required
def employee_detail(request, username):
    """Tek calisan detay sayfasi."""
    if not request.user.is_yonetici() and not request.user.is_takim_lideri():
        return redirect('dashboard:employee_panel')

    from accounts.models import CustomUser
    employee = get_object_or_404(CustomUser, username=username)

    # Son 7 gunluk istatistikler
    today = date.today()
    week_ago = today - timedelta(days=7)
    stats = DailyStats.objects.filter(
        user=employee, date__gte=week_ago
    ).order_by('date')

    # Son postur ve calisma loglari
    recent_sessions = EmployeeSession.objects.filter(user=employee).order_by('-started_at')[:5]

    return render(request, 'dashboard/employee_detail.html', {
        'employee': employee,
        'stats': stats,
        'recent_sessions': recent_sessions,
    })


@login_required
def logs_view(request):
    """Log goruntuleme ve filtreleme."""
    if not request.user.can_view_logs():
        return redirect('dashboard:employee_panel')

    # Filtreler
    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')
    username = request.GET.get('username', '')
    log_type = request.GET.get('type', 'posture')

    accessible_rooms = request.user.get_accessible_rooms()
    rooms = accessible_rooms

    if log_type == 'posture':
        logs = PostureLog.objects.filter(
            session__room__in=accessible_rooms
        ).select_related('session__user')

        if date_from:
            logs = logs.filter(timestamp__date__gte=date_from)
        if date_to:
            logs = logs.filter(timestamp__date__lte=date_to)
        if username:
            logs = logs.filter(session__user__username__icontains=username)

        logs = logs[:500]
    else:
        logs = WorkStatusLog.objects.filter(
            session__room__in=accessible_rooms
        ).select_related('session__user')

        if date_from:
            logs = logs.filter(timestamp__date__gte=date_from)
        if date_to:
            logs = logs.filter(timestamp__date__lte=date_to)
        if username:
            logs = logs.filter(session__user__username__icontains=username)

        logs = logs[:500]

    return render(request, 'dashboard/logs.html', {
        'logs': logs,
        'log_type': log_type,
        'date_from': date_from,
        'date_to': date_to,
        'username': username,
        'rooms': rooms,
    })


@login_required
def export_csv(request):
    """CSV indirme."""
    if not request.user.can_export():
        return HttpResponse(status=403)

    log_type = request.GET.get('type', 'posture')
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="officeguardian_{log_type}_{date.today()}.csv"'
    response.write('\ufeff')  # BOM for Excel UTF-8

    writer = csv.writer(response)

    if log_type == 'posture':
        writer.writerow(['Çalışan', 'Oda', 'Tarih', 'Boyun Açısı', 'Gövde Açısı', 'İyi Postür', 'Skor'])
        logs = PostureLog.objects.filter(
            session__room__created_by=request.user
        ).select_related('session__user', 'session__room')[:5000]

        for log in logs:
            writer.writerow([
                log.session.user.username,
                log.session.room.name,
                log.timestamp.strftime('%Y-%m-%d %H:%M:%S'),
                f'{log.neck_angle:.1f}',
                f'{log.torso_angle:.1f}',
                'Evet' if log.is_good else 'Hayır',
                log.score,
            ])
    else:
        writer.writerow(['Çalışan', 'Oda', 'Tarih', 'Çalışıyor', 'Yaw', 'Pitch'])
        logs = WorkStatusLog.objects.filter(
            session__room__created_by=request.user
        ).select_related('session__user', 'session__room')[:5000]

        for log in logs:
            writer.writerow([
                log.session.user.username,
                log.session.room.name,
                log.timestamp.strftime('%Y-%m-%d %H:%M:%S'),
                'Evet' if log.is_working else 'Hayır',
                f'{log.yaw:.1f}',
                f'{log.pitch:.1f}',
            ])

    return response


@login_required
def export_excel(request):
    """Excel (.xlsx) export - coklu sheet, formatli."""
    if not request.user.can_export():
        return HttpResponse(status=403)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # --- Renk ve stil tanimlari ---
    header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    green_fill = PatternFill(start_color='D1FAE5', end_color='D1FAE5', fill_type='solid')
    red_fill = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB'),
    )

    def style_header(ws, headers):
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border

    def auto_width(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
                cell.border = thin_border
            ws.column_dimensions[col_letter].width = min(max_len + 3, 30)

    # --- Sheet 1: Gunluk Istatistikler ---
    ws1 = wb.active
    ws1.title = 'Gunluk Istatistikler'
    headers1 = ['Calisan', 'Oda', 'Tarih', 'Calisma (sn)', 'Mola (sn)', 'Iyi Postur (sn)', 'Kotu Postur (sn)', 'Postur Skoru %', 'Verimlilik %']
    style_header(ws1, headers1)

    daily_stats = DailyStats.objects.filter(
        room__created_by=request.user
    ).select_related('user', 'room').order_by('-date')[:2000]

    for i, s in enumerate(daily_stats, 2):
        ws1.cell(row=i, column=1, value=s.user.username)
        ws1.cell(row=i, column=2, value=s.room.name)
        ws1.cell(row=i, column=3, value=str(s.date))
        ws1.cell(row=i, column=4, value=round(s.work_seconds))
        ws1.cell(row=i, column=5, value=round(s.break_seconds))
        ws1.cell(row=i, column=6, value=round(s.good_posture_seconds))
        ws1.cell(row=i, column=7, value=round(s.bad_posture_seconds))
        score_cell = ws1.cell(row=i, column=8, value=s.posture_score)
        score_cell.fill = green_fill if s.posture_score >= 70 else red_fill
        ratio_cell = ws1.cell(row=i, column=9, value=s.work_ratio)
        ratio_cell.fill = green_fill if s.work_ratio >= 70 else red_fill
    auto_width(ws1)

    # --- Sheet 2: Postur Loglari ---
    ws2 = wb.create_sheet('Postur Loglari')
    headers2 = ['Calisan', 'Oda', 'Tarih', 'Boyun Acisi', 'Govde Acisi', 'Iyi Postur', 'Skor']
    style_header(ws2, headers2)

    posture_logs = PostureLog.objects.filter(
        session__room__created_by=request.user
    ).select_related('session__user', 'session__room').order_by('-timestamp')[:5000]

    for i, log in enumerate(posture_logs, 2):
        ws2.cell(row=i, column=1, value=log.session.user.username)
        ws2.cell(row=i, column=2, value=log.session.room.name)
        ws2.cell(row=i, column=3, value=log.timestamp.strftime('%Y-%m-%d %H:%M:%S'))
        ws2.cell(row=i, column=4, value=round(log.neck_angle, 1))
        ws2.cell(row=i, column=5, value=round(log.torso_angle, 1))
        good_cell = ws2.cell(row=i, column=6, value='Evet' if log.is_good else 'Hayir')
        good_cell.fill = green_fill if log.is_good else red_fill
        ws2.cell(row=i, column=7, value=log.score)
    auto_width(ws2)

    # --- Sheet 3: Calisma Loglari ---
    ws3 = wb.create_sheet('Calisma Loglari')
    headers3 = ['Calisan', 'Oda', 'Tarih', 'Calisiyor', 'Yaw', 'Pitch']
    style_header(ws3, headers3)

    work_logs = WorkStatusLog.objects.filter(
        session__room__created_by=request.user
    ).select_related('session__user', 'session__room').order_by('-timestamp')[:5000]

    for i, log in enumerate(work_logs, 2):
        ws3.cell(row=i, column=1, value=log.session.user.username)
        ws3.cell(row=i, column=2, value=log.session.room.name)
        ws3.cell(row=i, column=3, value=log.timestamp.strftime('%Y-%m-%d %H:%M:%S'))
        work_cell = ws3.cell(row=i, column=4, value='Evet' if log.is_working else 'Hayir')
        work_cell.fill = green_fill if log.is_working else red_fill
        ws3.cell(row=i, column=5, value=round(log.yaw, 1))
        ws3.cell(row=i, column=6, value=round(log.pitch, 1))
    auto_width(ws3)

    # --- Sheet 4: Ozet ---
    ws4 = wb.create_sheet('Ozet')
    ws4.insert_rows(0)
    style_header(ws4, ['Metrik', 'Deger'])
    ws4.cell(row=2, column=1, value='Toplam Calisan Sayisi')
    ws4.cell(row=2, column=2, value=len(set(s.user_id for s in daily_stats)))
    ws4.cell(row=3, column=1, value='Toplam Oda Sayisi')
    ws4.cell(row=3, column=2, value=len(set(s.room_id for s in daily_stats)))
    ws4.cell(row=4, column=1, value='Toplam Gun Sayisi')
    ws4.cell(row=4, column=2, value=len(set(s.date for s in daily_stats)))
    if daily_stats:
        avg_posture = round(sum(s.posture_score for s in daily_stats) / len(daily_stats))
        avg_work = round(sum(s.work_ratio for s in daily_stats) / len(daily_stats))
        ws4.cell(row=5, column=1, value='Ortalama Postur Skoru')
        ws4.cell(row=5, column=2, value=f'%{avg_posture}')
        ws4.cell(row=6, column=1, value='Ortalama Verimlilik')
        ws4.cell(row=6, column=2, value=f'%{avg_work}')
    auto_width(ws4)

    # Response olustur
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="OfficeGuardian_Rapor_{date.today()}.xlsx"'
    return response


@login_required
def employee_panel(request):
    """Calisan kendi paneli."""
    if request.user.is_yonetici():
        return redirect('dashboard:admin_dashboard')

    today = date.today()
    week_ago = today - timedelta(days=7)

    # Gunluk istatistikler
    today_stats = DailyStats.objects.filter(user=request.user, date=today).first()

    # Son 7 gun
    weekly_stats = DailyStats.objects.filter(
        user=request.user, date__gte=week_ago
    ).order_by('date')

    # Aktif oturum
    active_session = EmployeeSession.objects.filter(
        user=request.user, is_online=True
    ).first()

    # Okunmamis uyarilar
    unread_warnings = []
    if active_session:
        unread_warnings = Warning.objects.filter(
            to_session=active_session, is_read=False
        ).order_by('-created_at')

    # Tum uyarilar (son 20)
    recent_warnings = Warning.objects.filter(
        to_session__user=request.user
    ).order_by('-created_at')[:20]

    return render(request, 'dashboard/employee_panel.html', {
        'today_stats': today_stats,
        'weekly_stats': weekly_stats,
        'active_session': active_session,
        'unread_warnings': unread_warnings,
        'recent_warnings': recent_warnings,
    })


@login_required
def mark_warning_read(request, warning_id):
    """Uyariyi okundu olarak isaretle."""
    warning = get_object_or_404(Warning, id=warning_id, to_session__user=request.user)
    warning.is_read = True
    warning.save()
    return JsonResponse({'status': 'ok'})


@login_required
def daily_report(request):
    """Gunluk rapor sayfasi - admin icin."""
    if not request.user.is_yonetici() and not request.user.is_takim_lideri():
        return redirect('dashboard:employee_panel')

    report_date = request.GET.get('date', '')
    if report_date:
        try:
            from datetime import datetime
            report_date = datetime.strptime(report_date, '%Y-%m-%d').date()
        except ValueError:
            report_date = date.today()
    else:
        report_date = date.today()

    rooms = request.user.get_accessible_rooms().filter(is_active=True)
    selected_room_code = request.GET.get('room', '')
    selected_room = None
    stats = []

    if selected_room_code:
        selected_room = rooms.filter(code=selected_room_code).first()

    if selected_room:
        stats = DailyStats.objects.filter(
            room=selected_room, date=report_date
        ).select_related('user').order_by('-work_seconds')
    elif rooms.exists():
        stats = DailyStats.objects.filter(
            room__in=rooms, date=report_date
        ).select_related('user', 'room').order_by('-work_seconds')

    # Ozet hesapla
    total_work = sum(s.work_seconds for s in stats)
    total_break = sum(s.break_seconds for s in stats)
    avg_posture = 0
    avg_work_ratio = 0
    if stats:
        avg_posture = round(sum(s.posture_score for s in stats) / len(stats))
        avg_work_ratio = round(sum(s.work_ratio for s in stats) / len(stats))

    return render(request, 'dashboard/daily_report.html', {
        'rooms': rooms,
        'selected_room': selected_room,
        'report_date': report_date,
        'stats': stats,
        'total_work': total_work,
        'total_break': total_break,
        'avg_posture': avg_posture,
        'avg_work_ratio': avg_work_ratio,
    })


@login_required
def user_manage(request):
    """Kullanici yonetimi - sadece super admin."""
    if not request.user.can_manage_users():
        return redirect('dashboard:index')

    from accounts.models import CustomUser

    if request.method == 'POST':
        action = request.POST.get('action')
        user_id = request.POST.get('user_id')

        if action == 'change_role' and user_id:
            new_role = request.POST.get('new_role', '')
            if new_role in dict(CustomUser.ROLE_CHOICES):
                CustomUser.objects.filter(id=user_id).exclude(id=request.user.id).update(role=new_role)

        elif action == 'assign_room' and user_id:
            room_id = request.POST.get('room_id')
            try:
                user = CustomUser.objects.get(id=user_id)
                room = Room.objects.get(id=room_id)
                user.assigned_rooms.add(room)
            except (CustomUser.DoesNotExist, Room.DoesNotExist):
                pass

        elif action == 'remove_room' and user_id:
            room_id = request.POST.get('room_id')
            try:
                user = CustomUser.objects.get(id=user_id)
                room = Room.objects.get(id=room_id)
                user.assigned_rooms.remove(room)
            except (CustomUser.DoesNotExist, Room.DoesNotExist):
                pass

        return redirect('dashboard:user_manage')

    users = CustomUser.objects.all().order_by('role', 'username')
    rooms = Room.objects.all()
    return render(request, 'dashboard/user_manage.html', {
        'users': users,
        'rooms': rooms,
        'role_choices': CustomUser.ROLE_CHOICES,
    })


@login_required
def notifications_api(request):
    """Bildirim listesi (AJAX)."""
    notifs = Notification.objects.filter(user=request.user).order_by('-created_at')[:30]
    unread_count = Notification.objects.filter(user=request.user, is_read=False).count()
    data = {
        'unread_count': unread_count,
        'notifications': [
            {
                'id': n.id,
                'type': n.notif_type,
                'title': n.title,
                'message': n.message,
                'is_read': n.is_read,
                'created_at': n.created_at.strftime('%d.%m.%Y %H:%M'),
                'room': n.related_room.name if n.related_room else None,
            }
            for n in notifs
        ],
    }
    return JsonResponse(data)


@require_POST
@login_required
def notifications_read(request):
    """Tum bildirimleri okundu isaretle (AJAX)."""
    Notification.objects.filter(user=request.user, is_read=False).update(is_read=True)
    return JsonResponse({'status': 'ok'})


@require_POST
@login_required
def notification_read_single(request, notif_id):
    """Tek bildirimi okundu isaretle."""
    Notification.objects.filter(id=notif_id, user=request.user).update(is_read=True)
    return JsonResponse({'status': 'ok'})


def api_personal_model(request, username):
    """Kisisel model verisini JSON dondur (client icin)."""
    try:
        pm = PersonalModel.objects.get(user__username=username)
        if pm.is_calibrated:
            return JsonResponse({
                'calibrated': True,
                'neck_threshold': round(pm.neck_threshold, 1),
                'torso_threshold': round(pm.torso_threshold, 1),
                'yaw_threshold': round(pm.yaw_threshold, 1),
                'pitch_threshold': round(pm.pitch_threshold, 1),
                'neck_angle_mean': round(pm.neck_angle_mean, 1),
                'torso_angle_mean': round(pm.torso_angle_mean, 1),
            })
    except PersonalModel.DoesNotExist:
        pass
    return JsonResponse({'calibrated': False})


@csrf_exempt
@require_POST
def api_calibration_save(request):
    """Kalibrasyon verilerini kaydet (client'tan gelen POST)."""
    import json as _json
    try:
        data = _json.loads(request.body)
        username = data.get('username', '')
        from accounts.models import CustomUser
        user = CustomUser.objects.get(username=username)

        pm, _ = PersonalModel.objects.get_or_create(user=user)
        pm.neck_angle_mean = data.get('neck_angle_mean', 0)
        pm.neck_angle_std = data.get('neck_angle_std', 5)
        pm.torso_angle_mean = data.get('torso_angle_mean', 0)
        pm.torso_angle_std = data.get('torso_angle_std', 3)
        pm.yaw_mean = data.get('yaw_mean', 0)
        pm.yaw_std = data.get('yaw_std', 5)
        pm.pitch_mean = data.get('pitch_mean', 0)
        pm.pitch_std = data.get('pitch_std', 5)
        pm.calibration_samples = data.get('samples', 0)
        pm.is_calibrated = True
        pm.calibrated_at = timezone.now()
        pm.calculate_thresholds()
        pm.save()

        return JsonResponse({
            'status': 'ok',
            'neck_threshold': round(pm.neck_threshold, 1),
            'torso_threshold': round(pm.torso_threshold, 1),
            'yaw_threshold': round(pm.yaw_threshold, 1),
            'pitch_threshold': round(pm.pitch_threshold, 1),
        })
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=400)


@csrf_exempt
@require_POST
def api_face_register(request):
    """Yuz encoding kaydı (client'tan gelen POST)."""
    import json as _json
    import base64
    try:
        data = _json.loads(request.body)
        username = data.get('username', '')
        encoding_b64 = data.get('encoding', '')

        from accounts.models import CustomUser
        user = CustomUser.objects.get(username=username)
        encoding_bytes = base64.b64decode(encoding_b64)

        fe, _ = FaceEncoding.objects.update_or_create(
            user=user,
            defaults={'encoding_data': encoding_bytes}
        )
        return JsonResponse({'status': 'ok'})
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=400)


def api_face_encodings(request, room_code):
    """Odadaki tum yuz encoding'lerini dondur (client icin)."""
    import base64
    room = Room.objects.filter(code=room_code).first()
    if not room:
        return JsonResponse({'encodings': []})

    # Odada oturumu olan kullanicilarin yuz verilerini getir
    sessions = EmployeeSession.objects.filter(room=room).select_related('user')
    user_ids = set(s.user_id for s in sessions)
    encodings = FaceEncoding.objects.filter(user_id__in=user_ids).select_related('user')

    data = []
    for fe in encodings:
        data.append({
            'username': fe.user.username,
            'encoding': base64.b64encode(fe.encoding_data).decode('utf-8'),
        })
    return JsonResponse({'encodings': data})


def get_active_rooms(request):
    """Aktif odalari JSON olarak dondur (Login gerektirmez, client icin)."""
    rooms = Room.objects.filter(is_active=True).values('name', 'code', 'created_by__username')
    return JsonResponse({'rooms': list(rooms)})


@login_required
def download_launcher(request, room_code):
    """Odaya ozel baslatici .bat dosyasi olustur ve indir."""
    room = get_object_or_404(Room, code=room_code)
    
    # Sunucu adresini dinamik olarak al (Host header'dan)
    # Eger yerel agda erisiliyorsa IP, localhost ise localhost gelir.
    host = request.get_host()
    
    # Kullanici adini al (Giris yapmis olan kullanici)
    username = request.user.username
    
    # Bat dosyasi icerigi (UTF--8 karakter sorunu olmamasi icin chcp 65001 ekleyebiliriz ama 
    # Python cagirdigi icin genelde sorun olmaz. Windows bat dosyalari genelde ANSI bekler 
    # ama Python parametrelerinde sorun cikabilir. Basit tutalim.)
    # 
    # Onemli: Musteri bilgisayarinda python yolunun PATH'de oldugunu varsayiyoruz.
    # Ayrica scriptin 'client/employee_client.py' yolunu bulabilmesi icin
    # kullanicinin projeyi indirdigi klasorde bu bat'i calistirmasi veya 
    # bat dosyasinin proje klasorunde olmasi gerekir.
    # En garantisi: Bat dosyasi projenin kok dizininde calistirilmalidir.
    
    bat_content = f"""@echo off
@chcp 65001 >nul
title OfficeGuardian - Odaya Baglaniyor...
echo ==============================================
echo  OfficeGuardian Baslatici
echo  Sunucu: {host}
echo  Oda: {room.name} ({room.code})
echo  Kullanici: {username}
echo ==============================================
echo.

:: Dosyanin bulundugu klasore git
cd /d "%~dp0"

:: Eger dosya 'OfficeGuardian' klasorunun icindeyse ama ana dizinde degilse (orn: indirilenler)
:: Bunu kontrol etmek zor, o yuzden kullaniciyi uyarmak en iyisi.

if not exist client\\employee_client.py (
    echo [HATA] 'client\\employee_client.py' dosyasi bulunamadi!
    echo.
    echo Lutfen bu dosyayi 'OfficeGuardian' proje klasorunun icine tasiyin.
    echo (manage.py dosyasinin oldugu yer)
    echo.
    pause
    exit
)

echo Istemci baslatiliyor...
echo.

:: Python komutunu calistir
python client/employee_client.py --server {host} --room {room.code} --user "{username}" --mode personal

if %errorlevel% neq 0 (
    echo.
    echo [HATA] Istemci hatali sonlandi veya Python bulunamadi.
    echo Lutfen Python'un yuklu oldugundan emin olun.
    echo.
    pause
)
"""
    
    # ZIP dosyasi olustur (Bellekte)
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        # Bat dosyasini zip icine ekle
        bat_filename = f"baglan_{room.code}_{username}.bat"
        zip_file.writestr(bat_filename, bat_content)

    # Pointer'i basa al
    zip_buffer.seek(0)

    # Response olustur
    response = HttpResponse(zip_buffer, content_type='application/zip')
    zip_filename = f"OfficeGuardian_Baglan_{username}.zip"
    response['Content-Disposition'] = f'attachment; filename="{zip_filename}"'
    return response


@login_required
def room_delete(request, room_id):
    """Odayi sil."""
    if not request.user.can_delete_room():
        return redirect('dashboard:employee_panel')

    if request.method == 'POST':
        room = get_object_or_404(Room, id=room_id)
        if not request.user.is_super_admin() and room.created_by != request.user:
            return redirect('dashboard:room_manage')
        room.delete()
        # Mesaj eklemek guzel olurdu ama simdilik redirect yeterli
        return redirect('dashboard:room_manage')
        
    return redirect('dashboard:room_manage')


def download_client_script(request):
    """Calisan icin zip indirme (Login gerektirmez)."""
    room_code = request.GET.get('room', '').strip()
    username = request.GET.get('user', '').strip()
    
    if not room_code or not username:
        return HttpResponse("Eksik parametre", status=400)
    
    room = get_object_or_404(Room, code=room_code)
    host = request.get_host()
    
    bat_content = f"""@echo off
@chcp 65001 >nul
title OfficeGuardian - Odaya Baglaniyor...
echo ==============================================
echo  OfficeGuardian Baslatici
echo  Sunucu: {host}
echo  Oda: {room.name} ({room.code})
echo  Kullanici: {username}
echo ==============================================
echo.

:: Dosyanin bulundugu klasore git
cd /d "%~dp0"

if not exist client\\employee_client.py (
    echo [HATA] 'client\\employee_client.py' dosyasi bulunamadi!
    echo.
    echo Lutfen bu dosyayi 'OfficeGuardian' proje klasorunun icine tasiyin.
    echo (manage.py dosyasinin oldugu yer)
    echo.
    pause
    exit
)

echo Istemci baslatiliyor...
echo.

:: Python komutunu calistir
python client/employee_client.py --server {host} --room {room.code} --user "{username}" --mode personal

if %errorlevel% neq 0 (
    echo.
    echo [HATA] Istemci hatali sonlandi veya Python bulunamadi.
    echo Lutfen Python'un yuklu oldugundan emin olun.
    echo.
    pause
)
"""
    # ZIP dosyasi olustur
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        bat_filename = f"baglan_{room.code}_{username}.bat"
        zip_file.writestr(bat_filename, bat_content)

    zip_buffer.seek(0)
    response = HttpResponse(zip_buffer, content_type='application/zip')
    zip_filename = f"OfficeGuardian_Baglan_{username}.zip"
    response['Content-Disposition'] = f'attachment; filename="{zip_filename}"'
    return response


def join_room(request):
    """Calisanin odaya katilmasi icin form."""
    if request.method == 'POST':
        room_code = request.POST.get('room_code', '').strip()
        username = request.POST.get('username', '').strip()
        
        if not room_code or not username:
             return render(request, 'dashboard/join_room.html', {
                'error': 'Lütfen tüm alanları doldurun.',
                'room_code': room_code,
                'username': username
            })

        room = Room.objects.filter(code=room_code).first()
        if not room:
            return render(request, 'dashboard/join_room.html', {
                'error': 'Geçersiz oda kodu!',
                'room_code': room_code,
                'username': username
            })
            
        # Komut stringini hazirla
        host = request.get_host()
        command = f'python client/employee_client.py --server {host} --room {room.code} --user "{username}"'
        
        # Basarili sayfasi goster (Komut ve Indirme butonu ile)
        return render(request, 'dashboard/join_room.html', {
            'success': True,
            'command': command,
            'room_code': room.code,
            'username': username,
        })

    return render(request, 'dashboard/join_room.html')
